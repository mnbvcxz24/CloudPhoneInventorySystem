import sqlite3
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import date
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

APP_NAME = "RFInventorySystem"
DB_PATH = Path(__file__).with_name("rf_inventory.db")
DEFAULT_RATE = 0.0035375
STATUSES = ("Available", "Reserved", "Claimed", "Expired")


class InventoryApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_NAME)
        self.root.geometry("1120x700")
        self.root.minsize(900, 580)
        self.db = sqlite3.connect(DB_PATH)
        self.db.row_factory = sqlite3.Row
        self.init_db()
        self.style()
        self.build_ui()
        self.refresh_all()

    def init_db(self):
        self.db.executescript("""
        PRAGMA foreign_keys = ON;
        CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL);
        INSERT OR IGNORE INTO settings(key,value) VALUES ('idr_to_php','0.0035375');
        CREATE TABLE IF NOT EXISTS products (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            type TEXT NOT NULL,
            days INTEGER NOT NULL DEFAULT 0,
            cost_idr REAL NOT NULL DEFAULT 0,
            selling_price REAL NOT NULL DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS codes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT NOT NULL UNIQUE,
            product_id INTEGER NOT NULL REFERENCES products(id),
            cost_idr REAL NOT NULL,
            cost_php REAL NOT NULL,
            status TEXT NOT NULL DEFAULT 'Available',
            date_claimed TEXT
        );
        """)
        self.db.commit()

    def rate(self):
        row = self.db.execute("SELECT value FROM settings WHERE key='idr_to_php'").fetchone()
        return float(row["value"]) if row else DEFAULT_RATE

    def style(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Treeview", rowheight=27, font=("Segoe UI", 10))
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        style.configure("TButton", padding=(9, 6))
        style.configure("Title.TLabel", font=("Segoe UI", 18, "bold"))
        style.configure("Metric.TLabel", font=("Segoe UI", 13, "bold"))

    def build_ui(self):
        header = ttk.Frame(self.root, padding=(16, 12))
        header.pack(fill="x")
        ttk.Label(header, text="RFInventorySystem", style="Title.TLabel").pack(side="left")
        ttk.Label(header, text="Redfinger code inventory • single-use tracking").pack(side="left", padx=18, pady=(7, 0))
        ttk.Button(header, text="Export to Excel", command=self.export_excel).pack(side="right")

        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill="both", expand=True, padx=14, pady=(0, 12))
        self.products_tab = ttk.Frame(self.notebook, padding=12)
        self.codes_tab = ttk.Frame(self.notebook, padding=12)
        self.summary_tab = ttk.Frame(self.notebook, padding=12)
        self.notebook.add(self.products_tab, text="Product & Sales")
        self.notebook.add(self.codes_tab, text="Code List")
        self.notebook.add(self.summary_tab, text="Profit Summary")
        self.build_products_tab()
        self.build_codes_tab()
        self.build_summary_tab()

    def build_products_tab(self):
        form = ttk.LabelFrame(self.products_tab, text="Add product plan", padding=10)
        form.pack(fill="x", pady=(0, 10))
        self.p_name = tk.StringVar()
        self.p_type = tk.StringVar(value="Standard")
        self.p_days = tk.StringVar(value="30")
        self.p_cost = tk.StringVar(value="0")
        self.p_price = tk.StringVar(value="0")
        fields = [("Name", self.p_name, 18), ("Type", self.p_type, 14), ("Days", self.p_days, 8),
                  ("Cost / code (IDR)", self.p_cost, 16), ("Selling price / code (PHP)", self.p_price, 18)]
        for i, (label, var, width) in enumerate(fields):
            ttk.Label(form, text=label).grid(row=0, column=i, sticky="w", padx=4)
            ttk.Entry(form, textvariable=var, width=width).grid(row=1, column=i, padx=4, pady=4)
        ttk.Button(form, text="Add plan", command=self.add_product).grid(row=1, column=len(fields), padx=8)
        cols = ("name", "type", "days", "stock", "cost_idr", "cost_php", "price", "sold_qty", "sold_amount")
        self.products_tree = ttk.Treeview(self.products_tab, columns=cols, show="headings")
        labels = ("Name", "Type", "Days", "Available stock", "Cost IDR", "Cost PHP", "Selling price", "Sold qty", "Sold amount")
        widths = (180, 100, 60, 110, 100, 100, 115, 75, 120)
        for col, label, width in zip(cols, labels, widths):
            self.products_tree.heading(col, text=label)
            self.products_tree.column(col, width=width, anchor="w" if col in ("name", "type") else "e")
        self.products_tree.pack(fill="both", expand=True)
        ttk.Label(self.products_tab, text="Available stock excludes Reserved, Claimed, and Expired codes.").pack(anchor="w", pady=(8, 0))

    def build_codes_tab(self):
        form = ttk.LabelFrame(self.codes_tab, text="Add single-use code", padding=10)
        form.pack(fill="x", pady=(0, 10))
        self.c_code = tk.StringVar()
        self.c_product = tk.StringVar()
        self.c_cost = tk.StringVar()
        self.c_status = tk.StringVar(value="Available")
        self.rate_var = tk.StringVar(value=str(self.rate()))
        ttk.Label(form, text="Code").grid(row=0, column=0, sticky="w", padx=4)
        ttk.Entry(form, textvariable=self.c_code, width=30).grid(row=1, column=0, padx=4, pady=4)
        ttk.Label(form, text="Product plan").grid(row=0, column=1, sticky="w", padx=4)
        self.product_combo = ttk.Combobox(form, textvariable=self.c_product, state="readonly", width=26)
        self.product_combo.grid(row=1, column=1, padx=4, pady=4)
        ttk.Label(form, text="Cost override (IDR, optional)").grid(row=0, column=2, sticky="w", padx=4)
        ttk.Entry(form, textvariable=self.c_cost, width=22).grid(row=1, column=2, padx=4, pady=4)
        ttk.Button(form, text="Add code", command=self.add_code).grid(row=1, column=3, padx=8)
        ttk.Label(form, text="IDR → PHP rate").grid(row=0, column=4, sticky="w", padx=4)
        ttk.Entry(form, textvariable=self.rate_var, width=13).grid(row=1, column=4, padx=4)
        ttk.Button(form, text="Save rate", command=self.save_rate).grid(row=1, column=5, padx=4)

        actions = ttk.Frame(self.codes_tab)
        actions.pack(fill="x", pady=(0, 8))
        ttk.Label(actions, text="Select a code, then change its status:").pack(side="left")
        self.status_combo = ttk.Combobox(actions, values=STATUSES, state="readonly", width=14)
        self.status_combo.set("Available")
        self.status_combo.pack(side="left", padx=8)
        ttk.Button(actions, text="Update status", command=self.update_status).pack(side="left")
        ttk.Button(actions, text="Delete selected", command=self.delete_code).pack(side="left", padx=8)

        cols = ("no", "code", "name", "type", "days", "cost_idr", "cost_php", "status", "date_claimed")
        self.codes_tree = ttk.Treeview(self.codes_tab, columns=cols, show="headings")
        labels = ("No", "Code", "Name", "Type", "Days", "Cost IDR", "Locked cost PHP", "Availability", "Date claimed")
        widths = (50, 260, 160, 100, 60, 100, 115, 110, 110)
        for col, label, width in zip(cols, labels, widths):
            self.codes_tree.heading(col, text=label)
            self.codes_tree.column(col, width=width, anchor="w" if col in ("code", "name", "type", "status", "date_claimed") else "e")
        self.codes_tree.pack(fill="both", expand=True)
        ttk.Label(self.codes_tab, text="Changing a code to Claimed records today's date and counts it as sold. Locked PHP costs never change when the rate changes.").pack(anchor="w", pady=(8, 0))

    def build_summary_tab(self):
        self.metric_revenue = tk.StringVar()
        self.metric_cost = tk.StringVar()
        self.metric_profit = tk.StringVar()
        row = ttk.Frame(self.summary_tab)
        row.pack(fill="x", pady=10)
        for title, var in (("Revenue", self.metric_revenue), ("Cost of claimed codes", self.metric_cost), ("Profit", self.metric_profit)):
            card = ttk.LabelFrame(row, text=title, padding=18)
            card.pack(side="left", fill="x", expand=True, padx=6)
            ttk.Label(card, textvariable=var, style="Metric.TLabel").pack(anchor="w")
        ttk.Label(self.summary_tab, text="Only codes marked Claimed count as sales. Reserved codes are not included in revenue or profit.").pack(anchor="w", padx=6, pady=8)
        ttk.Button(self.summary_tab, text="Refresh totals", command=self.refresh_all).pack(anchor="w", padx=6, pady=8)

    def add_product(self):
        try:
            name = self.p_name.get().strip()
            typ = self.p_type.get().strip() or "Standard"
            days = int(self.p_days.get())
            cost = float(self.p_cost.get())
            price = float(self.p_price.get())
            if not name or days < 0 or cost < 0 or price < 0:
                raise ValueError
            self.db.execute("INSERT INTO products(name,type,days,cost_idr,selling_price) VALUES(?,?,?,?,?)",
                            (name, typ, days, cost, price))
            self.db.commit()
            self.p_name.set("")
            self.refresh_all()
        except (ValueError, sqlite3.IntegrityError):
            messagebox.showerror("Invalid product", "Enter a name, whole number of days, and valid non-negative costs/prices.")

    def add_code(self):
        try:
            code = self.c_code.get().strip()
            product = self.selected_product_id()
            if not code or product is None:
                raise ValueError
            p = self.db.execute("SELECT cost_idr FROM products WHERE id=?", (product,)).fetchone()
            cost = float(self.c_cost.get()) if self.c_cost.get().strip() else float(p["cost_idr"])
            if cost < 0:
                raise ValueError
            locked_php = cost * self.rate()
            self.db.execute("INSERT INTO codes(code,product_id,cost_idr,cost_php,status) VALUES(?,?,?,?, 'Available')",
                            (code, product, cost, locked_php))
            self.db.commit()
            self.c_code.set("")
            self.c_cost.set("")
            self.refresh_all()
        except ValueError:
            messagebox.showerror("Invalid code", "Enter a unique code, choose a product, and use a valid non-negative cost.")
        except sqlite3.IntegrityError:
            messagebox.showerror("Duplicate code", "That code already exists. Each code must be unique.")

    def selected_product_id(self):
        value = self.c_product.get()
        if not value:
            return None
        try:
            return int(value.split(" | ", 1)[0])
        except ValueError:
            return None

    def save_rate(self):
        try:
            value = float(self.rate_var.get())
            if value <= 0:
                raise ValueError
            self.db.execute("INSERT INTO settings(key,value) VALUES('idr_to_php',?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (str(value),))
            self.db.commit()
            self.refresh_all()
            messagebox.showinfo("Exchange rate saved", "New codes will use the updated rate. Existing code costs stay locked.")
        except ValueError:
            messagebox.showerror("Invalid rate", "Enter a positive number.")

    def update_status(self):
        selected = self.codes_tree.selection()
        if not selected:
            messagebox.showwarning("Select a code", "Choose a code first.")
            return
        code_id = int(selected[0])
        status = self.status_combo.get()
        row = self.db.execute("SELECT status,date_claimed FROM codes WHERE id=?", (code_id,)).fetchone()
        if status == "Claimed" and row["status"] != "Claimed":
            claimed = date.today().isoformat()
        elif status == "Claimed":
            claimed = row["date_claimed"] or date.today().isoformat()
        else:
            claimed = None
        self.db.execute("UPDATE codes SET status=?,date_claimed=? WHERE id=?", (status, claimed, code_id))
        self.db.commit()
        self.refresh_all()

    def delete_code(self):
        selected = self.codes_tree.selection()
        if not selected:
            messagebox.showwarning("Select a code", "Choose a code first.")
            return
        if not messagebox.askyesno("Delete code", "Delete the selected code permanently?"):
            return
        self.db.execute("DELETE FROM codes WHERE id=?", (int(selected[0]),))
        self.db.commit()
        self.refresh_all()

    def refresh_all(self):
        self.refresh_products()
        self.refresh_codes()
        self.refresh_product_combo()
        claimed = self.db.execute("""
            SELECT COALESCE(SUM(p.selling_price),0) AS revenue,
                   COALESCE(SUM(c.cost_php),0) AS cost
            FROM codes c JOIN products p ON p.id=c.product_id WHERE c.status='Claimed'
        """).fetchone()
        revenue, cost = float(claimed["revenue"]), float(claimed["cost"])
        self.metric_revenue.set(f"₱{revenue:,.2f}")
        self.metric_cost.set(f"₱{cost:,.2f}")
        self.metric_profit.set(f"₱{revenue-cost:,.2f}")

    def refresh_products(self):
        for item in self.products_tree.get_children():
            self.products_tree.delete(item)
        rows = self.db.execute("""
        SELECT p.*,
          (SELECT COUNT(*) FROM codes c WHERE c.product_id=p.id AND c.status='Available') stock,
          (SELECT COUNT(*) FROM codes c WHERE c.product_id=p.id AND c.status='Claimed') sold
        FROM products p ORDER BY p.id
        """).fetchall()
        for p in rows:
            self.products_tree.insert("", "end", iid=str(p["id"]), values=(
                p["name"], p["type"], p["days"], p["stock"], f'{p["cost_idr"]:,.2f}',
                f'₱{p["cost_idr"]*self.rate():,.2f}', f'₱{p["selling_price"]:,.2f}',
                p["sold"], f'₱{p["sold"]*p["selling_price"]:,.2f}'))

    def refresh_product_combo(self):
        rows = self.db.execute("SELECT id,name,type,days FROM products ORDER BY name").fetchall()
        self.product_combo["values"] = [f'{p["id"]} | {p["name"]} | {p["type"]} | {p["days"]} days' for p in rows]
        if self.c_product.get() not in self.product_combo["values"]:
            self.c_product.set(self.product_combo["values"][0] if self.product_combo["values"] else "")

    def refresh_codes(self):
        for item in self.codes_tree.get_children():
            self.codes_tree.delete(item)
        rows = self.db.execute("""
        SELECT c.*,p.name,p.type,p.days FROM codes c JOIN products p ON p.id=c.product_id ORDER BY c.id
        """).fetchall()
        for i, c in enumerate(rows, 1):
            self.codes_tree.insert("", "end", iid=str(c["id"]), values=(
                i, c["code"], c["name"], c["type"], c["days"], f'{c["cost_idr"]:,.2f}',
                f'₱{c["cost_php"]:,.2f}', c["status"], c["date_claimed"] or ""))

    def export_excel(self):
        if Workbook is None:
            messagebox.showerror("Missing dependency", "Install dependencies with: pip install -r requirements.txt")
            return
        path = filedialog.asksaveasfilename(title="Export inventory", defaultextension=".xlsx",
            initialfile="RFInventorySystem.xlsx", filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "Product & Sales"
        ws.append(["Name", "Type", "Days", "Stock", "Cost in IDR", "Cost in PHP", "Selling Price", "Sold Qty", "Sold Amount"])
        products = self.db.execute("SELECT * FROM products ORDER BY id").fetchall()
        for p in products:
            codes = self.db.execute("SELECT status,cost_php FROM codes WHERE product_id=?", (p["id"],)).fetchall()
            available = sum(c["status"] == "Available" for c in codes)
            sold = sum(c["status"] == "Claimed" for c in codes)
            ws.append([p["name"], p["type"], p["days"], available, p["cost_idr"], p["cost_idr"]*self.rate(),
                       p["selling_price"], sold, sold*p["selling_price"]])
        cl = wb.create_sheet("Code List")
        cl.append(["No", "Code", "Name", "Type", "Days", "Availability", "Date Claimed", "Cost in IDR", "Locked Cost in PHP"])
        codes = self.db.execute("""SELECT c.*,p.name,p.type,p.days FROM codes c JOIN products p ON p.id=c.product_id ORDER BY c.id""").fetchall()
        for i, c in enumerate(codes, 1):
            cl.append([i, c["code"], c["name"], c["type"], c["days"], c["status"], c["date_claimed"] or "",
                       c["cost_idr"], c["cost_php"]])
        sm = wb.create_sheet("Profit Summary")
        claimed = [c for c in codes if c["status"] == "Claimed"]
        revenue = sum(float(self.db.execute("SELECT selling_price FROM products WHERE id=(SELECT product_id FROM codes WHERE id=?)", (c["id"],)).fetchone()[0]) for c in claimed)
        cost = sum(c["cost_php"] for c in claimed)
        sm.append(["Metric", "Amount (PHP)"])
        sm.append(["Revenue", revenue])
        sm.append(["Cost of claimed codes", cost])
        sm.append(["Profit", revenue-cost])
        for sheet in wb.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="183153")
                cell.alignment = Alignment(horizontal="center")
            for col in sheet.columns:
                letter = get_column_letter(col[0].column)
                width = min(max(max(len(str(cell.value or "")) for cell in col) + 2, 12), 32)
                sheet.column_dimensions[letter].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (int, float)) and cell.column > 1:
                        cell.number_format = '#,##0.00'
        try:
            wb.save(path)
            messagebox.showinfo("Export complete", f"Workbook saved to:\n{path}")
        except OSError as exc:
            messagebox.showerror("Export failed", str(exc))

    def close(self):
        self.db.close()
        self.root.destroy()


def main():
    root = tk.Tk()
    app = InventoryApp(root)
    root.protocol("WM_DELETE_WINDOW", app.close)
    root.mainloop()


if __name__ == "__main__":
    main()
